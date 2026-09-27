import os
import random
import requests
from datetime import datetime, timedelta, timezone
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import config
from db import execute_d1, init_d1_db

TURKEY_TZ = timezone(timedelta(hours=3))

def get_turkey_now():
    return datetime.now(TURKEY_TZ)

def telegram_post(text, chat_id=None, reply_markup=None):
    target_chat_id = chat_id or config.TELEGRAM_CHAT_ID
    if not target_chat_id or not config.TELEGRAM_BOT_TOKEN:
        print("❌ Telegram token veya Chat ID eksik.")
        return

    url = f"https://api.telegram.org/bot{config.TELEGRAM_BOT_TOKEN}/sendMessage"
    payload = {
        "chat_id": target_chat_id,
        "text": text,
        "parse_mode": "HTML",
        "disable_web_page_preview": True
    }
    if reply_markup:
        payload["reply_markup"] = reply_markup

    try:
        requests.post(url, json=payload, timeout=10)
    except Exception as e:
        print(f"❌ Telegram gönderim hatası: {e}")

def telegram_send_document(file_path, caption="", chat_id=None):
    target_chat_id = chat_id or config.TELEGRAM_CHAT_ID
    if not target_chat_id or not config.TELEGRAM_BOT_TOKEN:
        return

    url = f"https://api.telegram.org/bot{config.TELEGRAM_BOT_TOKEN}/sendDocument"
    try:
        with open(file_path, "rb") as f:
            files = {"document": f}
            data = {"chat_id": target_chat_id, "caption": caption, "parse_mode": "HTML"}
            requests.post(url, data=data, files=files, timeout=30)
    except Exception as e:
        print(f"❌ Telegram dosya gönderim hatası: {e}")

def excel_dosyasi_olustur_ve_gonder(chat_id=None):
    init_d1_db()
    maclar = execute_d1("SELECT * FROM maclar") or []

    if not maclar:
        telegram_post("📊 Veritabanında dışa aktarılacak oran analizi bulunamadı.", chat_id)
        return

    # Excel Workbook oluşturma
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Bahis Analizleri"

    # Tablo Başlıkları
    headers = ["ID", "Zaman / Saat", "Ev Sahibi / Maç", "Deplasman / Pazar", "Lig / Kategori", "Tahmin / Detay"]
    ws.append(headers)

    # Başlık Stili (Koyu Gri Zemin, Beyaz Kalın Harfler)
    header_fill = PatternFill(start_color="333333", end_color="333333", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    align_center = Alignment(horizontal="center", vertical="center")

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center

    ws.row_dimensions[1].height = 25

    # Verileri Yazdırma
    thin_border = Border(
        left=Side(style='thin', color='DDDDDD'),
        right=Side(style='thin', color='DDDDDD'),
        top=Side(style='thin', color='DDDDDD'),
        bottom=Side(style='thin', color='DDDDDD')
    )

    for idx, m in enumerate(maclar, start=2):
        row_data = [
            m.get("id", idx-1),
            m.get("saat", "Analiz"),
            m.get("ev_sahibi", ""),
            m.get("deplasman", ""),
            m.get("lig", ""),
            m.get("tahmin", "")
        ]
        ws.append(row_data)
        
        ws.row_dimensions[idx].height = 20
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=idx, column=col_num)
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="left", vertical="center")

    # Sütun Genişliklerini Otomatik Ayarlama
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 5, 15)

    # Gridlines Açık Tutma
    ws.views.sheetView[0].showGridLines = True

    # Dosyayı Kaydetme
    filename = f"Bahis_Analizleri_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    wb.save(filename)

    # Telegram üzerinden kullanıcıya dosya olarak gönderme
    caption = f"📈 <b>Güncel Bahis Oran Analizleri Raporu</b>\nToplam <b>{len(maclar)}</b> fırsat Excel olarak dışa aktarıldı."
    telegram_send_document(filename, caption, chat_id)

    # Gönderim sonrası yerel dosyayı temizleme
    try:
        os.remove(filename)
    except:
        pass

def bulteni_apiden_veritabanina_yukle(chat_id=None, dt_obj=None):
    init_d1_db()
    
    url = f"{config.BASE_URL}/v1/advantages/"

    headers = {
        "x-rapidapi-key": config.RAPIDAPI_KEY,
        "x-rapidapi-host": config.RAPIDAPI_HOST
    }

    try:
        print(f"🔍 Oran analizi isteği atılıyor -> URL: {url}")
        res = requests.get(url, headers=headers, timeout=25)
        
        status_code = res.status_code
        raw_text = res.text[:400]
        
        debug_info = f"🛠 <b>ORAN API DEBUG</b>\n- Status Code: <b>{status_code}</b>\n- Yanıt Özeti: <pre>{raw_text}</pre>"
        telegram_post(debug_info, chat_id)

        if status_code != 200:
            telegram_post(f"❌ Oran API Hata Döndürdü! Kod: {status_code}", chat_id)
            return

        data = res.json()
        
        items = []
        if isinstance(data, dict):
            advantages_obj = data.get("advantages", {})
            if isinstance(advantages_obj, dict):
                for key, val in advantages_obj.items():
                    if isinstance(val, list):
                        items.extend(val)
            elif isinstance(advantages_obj, list):
                items = advantages_obj
        elif isinstance(data, list):
            items = data

        if not items:
            telegram_post("⚠️ API'den yanıt alındı ancak advantages içeriği boş veya şemaya uymuyor.", chat_id)
            return

        yeni_eklenen = 0
        for item in items:
            if not isinstance(item, dict):
                continue
            
            val_type = item.get("type", "EV")
            val_oran = item.get("value", 0)
            val_type_str = item.get("valueType", "")
            market = item.get("marketKey", "Genel Pazar")
            event_id = item.get("eventKey", "Bilinmeyen Maç")
            
            lig = "Değerli Oran (Plus EV)"
            ev = f"ID: {event_id[:10]}..." if event_id else "Maç"
            dep = f"Pazar: {market[:15]}..." if market else "Pazar"
            detay = f"Tür: {val_type} | Değer: {val_oran} ({val_type_str})"

            execute_d1(
                "INSERT INTO maclar (saat, ev_sahibi, deplasman, lig, tahmin) VALUES (?, ?, ?, ?, ?)",
                ["Analiz", str(ev), str(dep), str(lig), str(detay)]
            )
            yeni_eklenen += 1

        telegram_post(f"✅ <b>Avantaj Analizleri Yüklendi!</b> ➕ <b>{yeni_eklenen}</b> fırsat veritabanına eklendi. Excel raporu hazırlanıyor...", chat_id)
        
        # Sayfalandırma kaldırıldı, doğrudan Excel dosyası oluşturulup gönderiliyor
        excel_dosyasi_olustur_ve_gonder(chat_id)

    except Exception as e:
        err_msg = f"❌ <b>Veri İşleme Kritik Hatası:</b>\n<code>{str(e)}</code>"
        print(err_msg)
        telegram_post(err_msg, chat_id)

def yarin_bultenini_yukle(chat_id=None):
    bulteni_apiden_veritabanina_yukle(chat_id=chat_id)
